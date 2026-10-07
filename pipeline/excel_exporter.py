import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from config import EXPORTS_DIR

class ExcelExporter:
    """Exports processed Zenlora product pipeline and UGC prompts into formatted Excel files."""

    # Zenlora Brand Colors
    HEADER_BG = "0D3B2E"      # Royal Emerald Green
    HEADER_TEXT = "FFFFFF"    # White
    GOLD_ACCENT = "E5C158"    # Metallic Gold
    ZEBRA_BG = "F7FAF8"       # Soft Mint Light Gray
    BORDER_COLOR = "D0DCD5"

    COLUMNS = [
        ("Zenlora Catchy Title", 28),
        ("Primary Category", 24),
        ("Content Bucket", 22),
        ("DeoDap Wholesale (₹)", 18),
        ("Zenlora Price (₹)", 16),
        ("MRP (₹)", 14),
        ("Gross Profit (₹)", 16),
        ("Margin (%)", 14),
        ("On-Screen Hook Text", 30),
        ("CTA Keyword", 14),
        ("Storyboard Image Prompt (9:16)", 40),
        ("Video Scene 1 (0-5s)", 35),
        ("Video Scene 2 (5-10s)", 40),
        ("Video Scene 3 (10-15s)", 40),
        ("Audio Recommendation", 28),
        ("DeoDap Source Link", 25),
        ("Image Preview Link", 25),
        ("SKU", 15)
    ]

    @classmethod
    def export(cls, products: list, filename_prefix: str = "Zenlora_Product_Pipeline") -> str:
        wb = Workbook()
        ws = wb.active
        ws.title = "Zenlora Pipeline"

        # Ensure grid lines are visible
        ws.views.sheetView[0].showGridLines = True

        # Define Styles
        font_header = Font(name="Calibri", size=11, bold=True, color=cls.HEADER_TEXT)
        fill_header = PatternFill(start_color=cls.HEADER_BG, end_color=cls.HEADER_BG, fill_type="solid")
        
        font_regular = Font(name="Calibri", size=10)
        font_bold = Font(name="Calibri", size=10, bold=True)
        font_gold = Font(name="Calibri", size=10, bold=True, color="996515")
        
        fill_zebra = PatternFill(start_color=cls.ZEBRA_BG, end_color=cls.ZEBRA_BG, fill_type="solid")
        fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        
        thin_border = Border(
            left=Side(style='thin', color=cls.BORDER_COLOR),
            right=Side(style='thin', color=cls.BORDER_COLOR),
            top=Side(style='thin', color=cls.BORDER_COLOR),
            bottom=Side(style='thin', color=cls.BORDER_COLOR)
        )

        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        align_right = Alignment(horizontal="right", vertical="center")

        # 1. Write Header Row
        for col_idx, (col_name, _) in enumerate(cls.COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = thin_border
            ws.row_dimensions[1].height = 28

        # 2. Write Data Rows
        for row_idx, p in enumerate(products, start=2):
            fill = fill_zebra if row_idx % 2 == 0 else fill_white
            ws.row_dimensions[row_idx].height = 45

            row_data = [
                p.get("zenlora_title", ""),
                p.get("primary_category", ""),
                p.get("content_bucket", ""),
                p.get("wholesale_price", 0.0),
                p.get("selling_price", 0),
                p.get("mrp", 0),
                p.get("gross_profit", 0.0),
                f"{p.get('profit_margin_pct', 0.0)}%",
                p.get("on_screen_hook_text", ""),
                p.get("cta_keyword", ""),
                p.get("storyboard_image_prompt", ""),
                p.get("scene_1_video_prompt", ""),
                p.get("scene_2_video_prompt", ""),
                p.get("scene_3_video_prompt", ""),
                p.get("audio_recommendation", ""),
                p.get("deodap_url", ""),
                p.get("primary_image", ""),
                p.get("sku", "")
            ]

            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = font_regular
                cell.fill = fill
                cell.border = thin_border

                # Formatting by column type
                if col_idx in [4, 5, 6, 7]:  # Pricing columns
                    cell.alignment = align_right
                    cell.number_format = '₹#,##0.00' if isinstance(val, float) else '₹#,##0'
                    if col_idx == 7:  # Gross Profit highlight
                        cell.font = font_gold
                elif col_idx in [2, 3, 8, 10]:  # Category, Bucket, Margin%, CTA
                    cell.alignment = align_center
                    if col_idx == 10:
                        cell.font = font_bold
                elif col_idx in [16, 17]:  # URLs
                    cell.alignment = align_left
                    if val and str(val).startswith("http"):
                        cell.hyperlink = str(val)
                        cell.font = Font(name="Calibri", size=10, underline="single", color="1B5E20")
                else:
                    cell.alignment = align_left

        # 3. Set Column Widths
        for col_idx, (_, width) in enumerate(cls.COLUMNS, start=1):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width

        # 4. Save Workbook
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{filename_prefix}_{timestamp}.xlsx"
        filepath = EXPORTS_DIR / filename
        wb.save(filepath)

        return str(filepath)
